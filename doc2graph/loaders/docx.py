"""
DOCX loader for doc2graph — extracts text from Microsoft Word documents.

Uses ``python-docx`` (MIT license).  Paragraphs are joined with newlines;
table cell text is extracted row-by-row, tab-separated.

Installation
------------
::

    pip install doc2graph[docx]
"""

from __future__ import annotations

from pathlib import Path
from typing import List


def _require_docx():
    try:
        import docx  # noqa: F401
        return docx
    except ImportError:
        raise ImportError(
            "python-docx is required for DOCX loading. "
            "Install it with: pip install doc2graph[docx]"
        )


def load_docx(path: str, include_tables: bool = True) -> str:
    """
    Extract text from a Microsoft Word (.docx) file.

    Paragraphs are returned in document order, separated by newlines.
    Table cells are extracted row-by-row with tab separation and appended
    after the main body paragraphs.

    Parameters
    ----------
    path : str
        Path to the ``.docx`` file.
    include_tables : bool
        Whether to extract text from tables.  Default ``True``.

    Returns
    -------
    str
        Full document text.  Empty paragraphs are omitted.

    Raises
    ------
    ImportError
        If ``python-docx`` is not installed.
    FileNotFoundError
        If the file does not exist.
    ValueError
        If the file extension is not ``.docx``.

    Example
    -------
    >>> from doc2graph.loaders.docx import load_docx
    >>> text = load_docx("report.docx")
    >>> text = load_docx("contract.docx", include_tables=False)
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {path}")
    if p.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file, got '{p.suffix}'")

    docx = _require_docx()
    doc = docx.Document(str(p))

    parts: List[str] = []

    # Body paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Tables
    if include_tables:
        for table in doc.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

    return "\n".join(parts)
