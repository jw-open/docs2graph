"""
PPTX loader for doc2graph — extracts text from PowerPoint presentations.

Uses ``python-pptx`` (MIT license).  Text is extracted slide-by-slide,
shape-by-shape, in presentation order.

Installation
------------
::

    pip install docs2graph[pptx]
"""

from __future__ import annotations

from pathlib import Path
from typing import List


def _require_pptx():
    try:
        from pptx import Presentation  # noqa: F401
        return Presentation
    except ImportError:
        raise ImportError(
            "python-pptx is required for PPTX loading. "
            "Install it with: pip install docs2graph[pptx]"
        )


def load_pptx(path: str, slide_markers: bool = True) -> str:
    """
    Extract text from a PowerPoint (.pptx) file.

    Text is collected from every text frame in each slide, in presentation
    order.  Optionally, slide boundary markers (``--- Slide N ---``) are
    inserted between slides so downstream models understand the structure.

    Parameters
    ----------
    path : str
        Path to the ``.pptx`` file.
    slide_markers : bool
        Insert ``--- Slide N ---`` markers between slides.  Default ``True``.

    Returns
    -------
    str
        Full presentation text.

    Raises
    ------
    ImportError
        If ``python-pptx`` is not installed.
    FileNotFoundError
        If the file does not exist.
    ValueError
        If the file extension is not ``.pptx``.

    Example
    -------
    >>> from docs2graph.loaders.pptx import load_pptx
    >>> text = load_pptx("deck.pptx")
    >>> text = load_pptx("slides.pptx", slide_markers=False)
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {path}")
    if p.suffix.lower() != ".pptx":
        raise ValueError(f"Expected a .pptx file, got '{p.suffix}'")

    Presentation = _require_pptx()
    prs = Presentation(str(p))

    slide_parts: List[str] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts: List[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if line:
                    texts.append(line)

        if texts:
            if slide_markers:
                slide_parts.append(f"--- Slide {slide_num} ---\n\n" + "\n".join(texts))
            else:
                slide_parts.append("\n".join(texts))

    return "\n\n".join(slide_parts)
